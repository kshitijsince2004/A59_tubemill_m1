"""Extract complete field dictionaries from A59 M1 data-mapping workbooks."""
from __future__ import annotations

from collections import Counter, defaultdict
from openpyxl import load_workbook
import json
import os

BASE = r"c:\dev\A59 Tube Mill M1\Dev Specs\SHEET"
FILES = {
    "TubeMill": "Zedral_A59_TubeMill_M1_Data_Mapping.xlsx",
    "Furnace": "Zedral_A59_Furnace_M1_Data_Mapping.xlsx",
    "STP": "Zedral_A59_STP_M1_Data_Mapping.xlsx",
    "DrawBench": "Zedral_A59_DrawBench_M1_Data_Mapping.xlsx",
}


def cell_fill_rgb(cell):
    try:
        fg = cell.fill.fgColor
        if fg is None:
            return None
        if fg.type == "rgb" and fg.rgb and fg.rgb not in ("00000000", "0"):
            return str(fg.rgb)
        if fg.type == "theme" and fg.theme is not None:
            return f"theme:{fg.theme}"
        if fg.type == "indexed" and fg.indexed is not None and fg.indexed != 64:
            return f"idx:{fg.indexed}"
    except Exception:
        pass
    return None


def sheet_to_rows(ws, header_row=3):
    headers = []
    for c in range(1, ws.max_column + 1):
        h = ws.cell(header_row, c).value
        headers.append(str(h).strip() if h is not None else f"col_{c}")
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
            continue
        first = vals[0]
        if isinstance(first, str) and first.strip().startswith("---"):
            continue
        row = {}
        for h, v in zip(headers, vals):
            if v is None:
                row[h] = None
            elif isinstance(v, str):
                row[h] = v.strip()
            else:
                row[h] = v
        class_idx = None
        for i, h in enumerate(headers):
            if h.lower() in ("class", "capture class"):
                class_idx = i + 1
                break
        if class_idx:
            row["_class_fill"] = cell_fill_rgb(ws.cell(r, class_idx))
        rows.append(row)
    return headers, rows


def cover_legend(ws):
    items = []
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        if a is None and b is None:
            continue
        items.append(
            {
                "key": a,
                "value": b,
                "fill": cell_fill_rgb(ws.cell(r, 1)) or cell_fill_rgb(ws.cell(r, 2)),
            }
        )
    return items


def dump_table(ws, header_row=3):
    headers, rows = sheet_to_rows(ws, header_row)
    return {"headers": headers, "rows": rows, "row_count": len(rows)}


def extract_all():
    results = {}
    for key, fname in FILES.items():
        path = os.path.join(BASE, fname)
        wb = load_workbook(path, data_only=True)
        wb_fmt = load_workbook(path, data_only=False)
        book = {
            "file": fname,
            "sheets": wb.sheetnames,
            "cover_legend": cover_legend(wb_fmt["Cover & Legend"])
            if "Cover & Legend" in wb.sheetnames
            else [],
            "sheet_map": dump_table(wb["Sheet Map"]) if "Sheet Map" in wb.sheetnames else None,
            "field_register": None,
            "class_counts": {},
            "by_target_table": {},
            "by_sheet_group": {},
            "erp_wo_fields": dump_table(wb["ERP-WO Fields"])
            if "ERP-WO Fields" in wb.sheetnames
            else None,
            "codes_enums": dump_table(wb["Codes & Enums"])
            if "Codes & Enums" in wb.sheetnames
            else None,
            "data_flow": dump_table(wb["Data Flow"]) if "Data Flow" in wb.sheetnames else None,
            "required_documents": dump_table(wb["Required Documents"])
            if "Required Documents" in wb.sheetnames
            else None,
            "extra_sheets": {},
            "class_fill_map": {},
        }
        if "Field Register" in wb.sheetnames:
            headers, rows = sheet_to_rows(wb_fmt["Field Register"], 3)
            _, rows_data = sheet_to_rows(wb["Field Register"], 3)
            for i, rd in enumerate(rows_data):
                if i < len(rows):
                    for h in headers:
                        if rd.get(h) is not None:
                            rows[i][h] = rd[h]
            book["field_register"] = {"headers": headers, "rows": rows, "row_count": len(rows)}
            classes = Counter()
            by_table = defaultdict(list)
            by_group = defaultdict(list)
            fills = defaultdict(set)
            for row in rows:
                cls = row.get("Class") or "UNKNOWN"
                classes[cls] += 1
                fills[cls].add(row.get("_class_fill"))
                tbl = row.get("Target schema.table") or "(none)"
                by_table[tbl].append(row.get("Field ID") or row.get("Canonical field (DB)"))
                grp = row.get("Sheet") or "(none)"
                by_group[grp].append(row.get("Field ID") or row.get("Canonical field (DB)"))
            book["class_counts"] = dict(classes)
            book["by_target_table"] = dict(by_table)
            book["by_sheet_group"] = dict(by_group)
            book["class_fill_map"] = {k: sorted(v, key=lambda x: str(x)) for k, v in fills.items()}

        for ename in wb.sheetnames:
            if ename in (
                "Cover & Legend",
                "Sheet Map",
                "Field Register",
                "ERP-WO Fields",
                "Codes & Enums",
                "Data Flow",
                "Required Documents",
            ):
                continue
            book["extra_sheets"][ename] = dump_table(wb[ename])

        results[key] = book
        print(f"{key}: {book['field_register']['row_count']} fields | {book['class_counts']}")
    return results


def esc(v):
    if v is None:
        return ""
    return str(v).replace("|", "\\|").replace("\n", " ")


def write_markdown(data):
    lines = []
    lines.append("# Zedral A59 M1 — Complete Field Dictionaries")
    lines.append("")
    lines.append(
        "Parsed from `Dev Specs/SHEET` data-mapping workbooks. "
        "**Field Register** is the source of truth for DB/forms planning."
    )
    lines.append("")
    lines.append("## Colour coding (capture Class column fills)")
    lines.append("")
    lines.append("Observed across workbooks from cell fills on the Class column:")
    lines.append("")
    lines.append("| Class | Typical fill | Meaning (from Cover & Legend) |")
    lines.append("| --- | --- | --- |")
    lines.append("| SYSTEM | `#EADCF2` lavender | System-generated id / timestamp |")
    lines.append("| DERIVED | `#D6E3F0` light blue | From ERP/WO, recipe chart, or computed |")
    lines.append("| AUTO | `#E2EFDA` / green-tint (when set) | PLC / SCADA / collector |")
    lines.append("| MANUAL | `#FFF2CC` / amber-tint (when set) | Operator / tablet entry |")
    lines.append("| AUTO/MANUAL | mixed | Prefer AUTO; fall back to MANUAL |")
    lines.append("| MASTER / MASTER/MANUAL | — | Master chart lookup (+ optional confirm) |")
    lines.append("")

    for key in ["TubeMill", "Furnace", "STP", "DrawBench"]:
        book = data[key]
        lines.append("---")
        lines.append("")
        lines.append(f"# {key}")
        lines.append("")
        lines.append(f"**File:** `{book['file']}`")
        sheet_list = ", ".join(f"`{s}`" for s in book["sheets"])
        lines.append(f"**Sheets ({len(book['sheets'])}):** {sheet_list}")
        lines.append(f"**Field Register count:** **{book['field_register']['row_count']}**")
        lines.append("")

        lines.append("## Capture class counts")
        lines.append("")
        for cls, n in sorted(book["class_counts"].items(), key=lambda x: (-x[1], x[0])):
            fill = book.get("class_fill_map", {}).get(cls, [])
            fill_s = ", ".join(str(x) for x in fill if x)
            extra = f" _(fill: {fill_s})_" if fill_s else ""
            lines.append(f"- **{cls}**: {n}{extra}")
        lines.append("")

        lines.append("## Target tables / child entities")
        lines.append("")
        for tbl, ids in book["by_target_table"].items():
            lines.append(f"- `{tbl}` — **{len(ids)}** fields: {', '.join(ids)}")
        lines.append("")

        lines.append("## Sheet / form groups")
        lines.append("")
        for grp, ids in book["by_sheet_group"].items():
            lines.append(f"- **{grp}** — {len(ids)}: {', '.join(ids)}")
        lines.append("")

        lines.append("## Cover & Legend")
        lines.append("")
        for item in book["cover_legend"]:
            k = item.get("key") or ""
            v = item.get("value") or ""
            if v:
                lines.append(f"- **{k}**: {v}")
            else:
                lines.append(f"- **{k}**")
        lines.append("")

        lines.append("## Sheet Map")
        lines.append("")
        if book["sheet_map"]:
            hdrs = [h for h in book["sheet_map"]["headers"] if not h.startswith("_")]
            lines.append("| " + " | ".join(hdrs) + " |")
            lines.append("| " + " | ".join(["---"] * len(hdrs)) + " |")
            for row in book["sheet_map"]["rows"]:
                lines.append("| " + " | ".join(esc(row.get(h)) for h in hdrs) + " |")
        lines.append("")

        lines.append("## COMPLETE Field Register")
        lines.append("")
        use_hdrs = [h for h in book["field_register"]["headers"] if not h.startswith("_")]
        lines.append("| " + " | ".join(use_hdrs) + " |")
        lines.append("| " + " | ".join(["---"] * len(use_hdrs)) + " |")
        for row in book["field_register"]["rows"]:
            lines.append("| " + " | ".join(esc(row.get(h)) for h in use_hdrs) + " |")
        lines.append("")

        if book.get("erp_wo_fields"):
            lines.append("## ERP / WO Fields")
            lines.append("")
            eh = [h for h in book["erp_wo_fields"]["headers"] if not h.startswith("_")]
            lines.append("| " + " | ".join(eh) + " |")
            lines.append("| " + " | ".join(["---"] * len(eh)) + " |")
            for row in book["erp_wo_fields"]["rows"]:
                lines.append("| " + " | ".join(esc(row.get(h)) for h in eh) + " |")
            lines.append("")

        if book.get("codes_enums"):
            lines.append("## Codes & Enums")
            lines.append("")
            eh = [h for h in book["codes_enums"]["headers"] if not h.startswith("_")]
            lines.append("| " + " | ".join(eh) + " |")
            lines.append("| " + " | ".join(["---"] * len(eh)) + " |")
            for row in book["codes_enums"]["rows"]:
                lines.append("| " + " | ".join(esc(row.get(h)) for h in eh) + " |")
            lines.append("")

        if book.get("data_flow"):
            lines.append("## Data Flow")
            lines.append("")
            eh = [h for h in book["data_flow"]["headers"] if not h.startswith("_")]
            lines.append("| " + " | ".join(eh) + " |")
            lines.append("| " + " | ".join(["---"] * len(eh)) + " |")
            for row in book["data_flow"]["rows"]:
                lines.append("| " + " | ".join(esc(row.get(h)) for h in eh) + " |")
            lines.append("")

        if book.get("required_documents"):
            lines.append("## Required Documents")
            lines.append("")
            eh = [h for h in book["required_documents"]["headers"] if not h.startswith("_")]
            lines.append("| " + " | ".join(eh) + " |")
            lines.append("| " + " | ".join(["---"] * len(eh)) + " |")
            for row in book["required_documents"]["rows"]:
                lines.append("| " + " | ".join(esc(row.get(h)) for h in eh) + " |")
            lines.append("")

        if book.get("extra_sheets"):
            lines.append("## Extra / supporting sheets")
            lines.append("")
            for ename, edata in book["extra_sheets"].items():
                lines.append(f"### {ename} ({edata['row_count']} rows)")
                lines.append("")
                eh = [h for h in edata["headers"] if not h.startswith("_")]
                # Always dump full content for planning — user asked for complete lists
                # but large param charts stay as tables
                lines.append("| " + " | ".join(eh) + " |")
                lines.append("| " + " | ".join(["---"] * len(eh)) + " |")
                for row in edata["rows"]:
                    lines.append("| " + " | ".join(esc(row.get(h)) for h in eh) + " |")
                lines.append("")

    path = os.path.join(BASE, "_parsed_field_dictionaries.md")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print("Wrote", path, "lines=", len(lines))
    return path


def write_compact_json(data):
    """Cleaner JSON without _class_fill noise for planning."""
    clean = {}
    for key, book in data.items():
        fr_rows = []
        for row in book["field_register"]["rows"]:
            fr_rows.append({k: v for k, v in row.items() if not k.startswith("_")})
        clean[key] = {
            "file": book["file"],
            "sheets": book["sheets"],
            "field_count": book["field_register"]["row_count"],
            "class_counts": book["class_counts"],
            "class_fill_map": book["class_fill_map"],
            "by_target_table": book["by_target_table"],
            "by_sheet_group": book["by_sheet_group"],
            "cover_legend": book["cover_legend"],
            "sheet_map": book["sheet_map"],
            "fields": fr_rows,
            "erp_wo_fields": book["erp_wo_fields"],
            "codes_enums": book["codes_enums"],
            "data_flow": book["data_flow"],
            "required_documents": book["required_documents"],
            "extra_sheets": book["extra_sheets"],
        }
    path = os.path.join(BASE, "_parsed_field_dictionaries.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(clean, f, indent=2, ensure_ascii=False, default=str)
    print("Wrote", path)
    return path


if __name__ == "__main__":
    data = extract_all()
    write_compact_json(data)
    write_markdown(data)

    # Also print full field lists to stdout for chat response assembly
    for key in ["TubeMill", "Furnace", "STP", "DrawBench"]:
        book = data[key]
        print("\n" + "=" * 80)
        print(f"{key} FIELD REGISTER ({book['field_register']['row_count']})")
        print("=" * 80)
        for row in book["field_register"]["rows"]:
            fid = row.get("Field ID")
            paper = row.get("Paper field")
            canon = row.get("Canonical field (DB)")
            typ = row.get("Type")
            unit = row.get("Unit")
            cls = row.get("Class")
            src = row.get("Source")
            req = row.get("Req")
            tbl = row.get("Target schema.table")
            sheet = row.get("Sheet")
            notes = row.get("Notes") or ""
            val = row.get("Validation / rule") or ""
            print(
                f"{fid}\t{sheet}\t{paper}\t{canon}\t{tbl}\t{typ}\t{unit}\t{cls}\t{src}\t{req}\t{val}\t{notes}"
            )
